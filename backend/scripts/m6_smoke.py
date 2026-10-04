"""M6 端到端验证：周期监控 -> 变更对比 -> 告警推送 -> 导出。

会起一个**本地 webhook 接收器**来真的验证告警是否发出（而不是只看返回值）。

用法::

    $env:RECON_TOKEN="..."; python scripts/m6_smoke.py http://127.0.0.1:8794
"""

from __future__ import annotations

import json
import os
import sys
import threading
import time
import urllib.error
import urllib.request
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8794"
TOKEN = os.environ.get("RECON_TOKEN", "")
HOOK_PORT = 8795

received: list[dict] = []


class HookHandler(BaseHTTPRequestHandler):
    def do_POST(self):  # noqa: N802
        length = int(self.headers.get("Content-Length") or 0)
        raw = self.rfile.read(length) if length else b""
        try:
            received.append(json.loads(raw.decode("utf-8")))
        except Exception:
            received.append({"raw": raw.decode("utf-8", "replace")})
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(b'{"ok":true}')

    def log_message(self, *args):
        pass


def call(path: str, method: str = "GET", body=None, raw: bool = False):
    headers = {"Content-Type": "application/json"}
    if TOKEN:
        headers["Authorization"] = f"Bearer {TOKEN}"
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(BASE + path, data=data, method=method, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            payload = resp.read()
            if raw:
                return resp.status, payload, dict(resp.headers)
            text = payload.decode("utf-8", "replace")
            return resp.status, (json.loads(text) if text.strip()[:1] in "{[" else text)
    except urllib.error.HTTPError as e:
        return (e.code, e.read().decode("utf-8", "replace"), {}) if raw else (
            e.code, e.read().decode("utf-8", "replace"))


def wait_scan(scan_id: int, timeout: float = 120.0) -> dict:
    deadline = time.time() + timeout
    current: dict = {}
    while time.time() < deadline:
        _, current = call(f"/api/scans/{scan_id}")
        if isinstance(current, dict) and current.get("status") not in ("running", "finalizing"):
            return current
        time.sleep(0.5)
    return current


def main() -> int:
    hook = ThreadingHTTPServer(("127.0.0.1", HOOK_PORT), HookHandler)
    threading.Thread(target=hook.serve_forever, daemon=True).start()

    preset_path = Path(".testtmp/m6_offline.yml").resolve()
    preset_path.parent.mkdir(parents=True, exist_ok=True)
    preset_path.write_text(
        "name: m6offline\n"
        "description: 只跑离线演示源\n"
        "include:\n  - demo_expand\n"
        "settings:\n  max_events: 500\n",
        encoding="utf-8",
    )

    print("=== 配置：白名单 + 本地 webhook 告警 ===")
    status, _ = call("/api/settings", "PUT", {
        "authorized_targets": ["example.com"],
        "require_authorization": True,
        "max_concurrent_scans": 2,
        "notify": {
            "enabled": True,
            "webhook_url": f"http://127.0.0.1:{HOOK_PORT}/hook",
            "webhook_token": "hook-token",
            "min_changes": 1,
        },
    })
    print(f"  PUT /api/settings -> {status}")

    print("\n=== 创建监控（首次只建基线）===")
    status, monitor = call("/api/monitors", "POST", {
        "name": "m6-smoke",
        "targets": ["example.com"],
        "preset": str(preset_path),
        "overrides": ["modules.demo_expand.prefixes=www,api"],
        "interval_minutes": 60,
        "enabled": True,
        "notify": True,
    })
    print(f"  POST /api/monitors -> {status} id={monitor.get('id') if isinstance(monitor, dict) else monitor}")
    if status != 201:
        return 1
    monitor_id = monitor["id"]

    status, run1 = call(f"/api/monitors/{monitor_id}/run", "POST")
    print(f"  立即执行 -> {status} scan_id={run1.get('scan_id')}")
    final1 = wait_scan(run1["scan_id"])
    print(f"  第一次扫描: {final1.get('status')} 耗时 {final1.get('elapsed')}s")

    _, assets1 = call(f"/api/scans/{run1['scan_id']}/assets")
    print(f"  基线资产: 域名 {sorted(d['name'] for d in assets1['domains'])}")
    print(f"  此时收到的告警数: {len(received)}（应为 0 —— 首次不告警）")

    print("\n=== 改掉配置再跑一次，制造变化 ===")
    call(f"/api/monitors/{monitor_id}", "PUT", {
        "name": "m6-smoke",
        "targets": ["example.com"],
        "preset": str(preset_path),
        "overrides": ["modules.demo_expand.prefixes=www,dev"],
        "interval_minutes": 60,
        "enabled": True,
        "notify": True,
    })
    status, run2 = call(f"/api/monitors/{monitor_id}/run", "POST")
    print(f"  立即执行 -> {status} scan_id={run2.get('scan_id')}")
    final2 = wait_scan(run2["scan_id"])
    print(f"  第二次扫描: {final2.get('status')}")

    time.sleep(1.0)
    print(f"  此时收到的告警数: {len(received)}（应为 1）")
    if received:
        body = received[-1]
        text = body.get("text", "")
        print(f"  告警标题: {body.get('title')}")
        for line in text.splitlines():
            if line.startswith("-") or "所有" in line or "**" in line:
                print(f"    {line}")

    print("\n=== 变更接口 ===")
    _, diff = call(f"/api/scans/{run2['scan_id']}/diff")
    print(f"  /diff -> 对比 #{diff.get('old_scan_id')} → #{diff.get('new_scan_id')} 共 {diff.get('total')} 处")
    print(f"  counts: {json.dumps(diff.get('counts'), ensure_ascii=False)}")
    _, changes = call("/api/changes?limit=5")
    print(f"  /api/changes -> {len(changes)} 条记录, 最新 total={changes[0]['total'] if changes else '-'}")

    print("\n=== 导出 ===")
    for fmt, fname in (("json", "x.json"), ("csv", "x.csv"), ("xlsx", "x.xlsx")):
        status, payload, headers = call(
            f"/api/scans/{run2['scan_id']}/export?format={fmt}", raw=True
        )
        out = Path(".testtmp") / fname
        out.write_bytes(payload)
        print(f"  {fmt:5} -> HTTP {status} {len(payload):>7} 字节  {headers.get('Content-Type') or headers.get('content-type')}")

    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO((Path(".testtmp") / "x.xlsx").read_bytes()))
    print(f"  xlsx 工作表: {workbook.sheetnames}")
    print(f"  域名表行数: {workbook['域名'].max_row}")

    json_doc = json.loads((Path(".testtmp") / "x.json").read_text(encoding="utf-8"))
    print(f"  json 顶层键: {sorted(json_doc.keys())}, 资产键: {sorted(json_doc['assets'].keys())}")

    print("\n=== 清理 ===")
    print(f"  DELETE /api/monitors/{monitor_id} -> {call(f'/api/monitors/{monitor_id}', 'DELETE')[0]}")
    hook.shutdown()
    return 0


if __name__ == "__main__":
    sys.exit(main())
